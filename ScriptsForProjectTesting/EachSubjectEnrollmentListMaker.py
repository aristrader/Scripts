"""
Script: Subject Enrollment List Generator

Description:
------------
This script fetches subject-wise enrollment data from an API, processes the data, and saves it to individual Excel files for each subject option. 
These are the lists that are required to be sent to teachers.

Steps:
------
1. Define the course details ID (`course_details_id`) to fetch subjects and enrollments.

"""

import os
import pandas as pd
import logging
import requests
from openpyxl import load_workbook
from openpyxl.worksheet.protection import SheetProtection

# Set up logging
log_file_path = '/Users/swapnilagarwal/Visual_Studio_Projects/Results/TestingExcels/subject_enrollment_list_failure.log'
logging.basicConfig(filename=log_file_path, level=logging.ERROR,
                    format='%(asctime)s %(levelname)s: %(message)s')

# Set up logging for summary details
summary_log_file_path = '/Users/swapnilagarwal/Visual_Studio_Projects/Results/TestingExcels/subject_summary.log'

# Create a logger specifically for the summary log
summary_logging = logging.getLogger('summary_logger')
summary_logging.setLevel(logging.INFO)

# File handler for writing summary logs to file
summary_handler = logging.FileHandler(summary_log_file_path)
summary_handler.setLevel(logging.INFO)

# Formatter for the log file entries
summary_formatter = logging.Formatter('%(asctime)s - %(message)s')
summary_handler.setFormatter(summary_formatter)

# Add handler to the summary logger
summary_logging.addHandler(summary_handler)
summary_logging.propagate = False

# API URLs
BASE_URL = "http://localhost:8080/v2"
GET_SUBJECTS_URL = f"{BASE_URL}/subject-details/course/{{course_details_id}}"
GET_ENROLLMENTS_URL = f"{BASE_URL}/enrollment-and-marks/getBySubjectDetailsId/{{subject_details_id}}"
GET_STUDENT_URL = f"{BASE_URL}/students/get/id/{{student_id}}"

# Directory for output
output_directory = '/Users/swapnilagarwal/Visual_Studio_Projects/Results/TestingExcels/ListsSubjectWise'

# Step 1: Replace with actual course details ID
course_details_id = 1

def protect_excel_columns_by_header(file_path, unlock_columns_headers):
    """
    Function to protect all columns except specified unlock_columns based on column headers in an Excel file.
    """
    # Load the workbook and select the active worksheet
    workbook = load_workbook(file_path)
    sheet = workbook.active
    
    # Create a dictionary to map column headings to their index
    header_to_column_index = {cell.value: cell.column for cell in sheet[1]}  # First row (header) is used for mapping

    # Lock all cells by default
    for row in sheet.iter_rows():
        for cell in row:
            # Copy the protection state and set it to locked
            cell.protection = cell.protection.copy(locked=True)

    # Unlock specified columns based on the header name
    for row in sheet.iter_rows(min_row=2):  # Start from the second row to skip the header
        for header in unlock_columns_headers:
            column_index = header_to_column_index.get(header)
            if column_index:
                # Unlock the cell in the current row and that column index
                row[column_index - 1].protection = row[column_index - 1].protection.copy(locked=False)

    # Lock the header row (first row)
    for cell in sheet[1]:
        cell.protection = cell.protection.copy(locked=True)  # Lock all header cells

    # Enable worksheet protection
    sheet.protection.set_password('frdyutfvht6675874567ddfz')  # Optional: Add password protection if desired
    sheet.protection.enable()

    # Save the workbook
    workbook.save(file_path)

# Fetch all subjects
try:
    subjects_response = requests.get(GET_SUBJECTS_URL.format(course_details_id=course_details_id))
    subjects_response.raise_for_status()
    subjects = subjects_response.json().get('data', [])
except requests.exceptions.RequestException as e:
    logging.error(f"Failed to fetch subjects for course_details_id {course_details_id}: {str(e)}")
    raise SystemExit(f"Critical error: Unable to proceed due to API failure: {str(e)}")

# Group subjects by subject type
subjects_by_type = {}
for subject in subjects:
    subject_type = subject['subjectType']
    options_name = subject['optionsName']
    subject_details_id = subject['subjectDetailsId']
    
    if subject_type not in subjects_by_type:
        subjects_by_type[subject_type] = {}
    
    subjects_by_type[subject_type][options_name] = subject_details_id

# Iterate over each subject type and each option
for subject_type, options in subjects_by_type.items():
    subject_type_dir = os.path.join(output_directory, subject_type)
    os.makedirs(subject_type_dir, exist_ok=True)
    
    summary_data = {}
    
    for options_name, subject_details_id in options.items():
        try:
            enrollments_response = requests.get(GET_ENROLLMENTS_URL.format(subject_details_id=subject_details_id))
            enrollments_response.raise_for_status()
            enrollments = enrollments_response.json().get('data', [])
        except requests.exceptions.RequestException as e:
            logging.error(f"Failed to fetch enrollments for subject_details_id {subject_details_id}: {str(e)}")
            continue

        students_data = []
        serial_no = 1
        for enrollment in enrollments:
            student_id = enrollment['studentId']
            
            try:
                student_response = requests.get(GET_STUDENT_URL.format(student_id=student_id))
                student_response.raise_for_status()
                student = student_response.json().get('data', {})
                
                student_info = {
                    'subject_id': subject_details_id,
                    'student_id': student_id,
                    'S.No': serial_no,
                    'roll_no': student['rollNo'],
                    'enrollment_no': student['enrollmentNo'],
                    'student_name': student['studentName'],
                    'main_marks': '',  # Empty column for main marks
                    'cce': '',         # Empty column for CCE
                    'practical_marks': ''  # Empty column for practical marks
                }
                students_data.append(student_info)
                serial_no += 1
            
            except requests.exceptions.RequestException as e:
                logging.error(f"Failed to fetch student details for student_id {student_id}: {str(e)}")
                continue

        df = pd.DataFrame(students_data)
        file_path = os.path.join(subject_type_dir, f"{options_name}.xlsx")
        
        if not df.empty:
            df.to_excel(file_path, index=False, sheet_name='sheet_1')
            print(f"Saved enrollment data for {options_name} under {subject_type} at {file_path}")
            
            summary_data[options_name] = len(students_data)
            
            # Lock all columns except "main_marks", "cce", and "practical_marks" using column headers
            protect_excel_columns_by_header(file_path, unlock_columns_headers=['main_marks', 'cce', 'practical_marks'])
            
            # Now, set column widths and hide the first two columns
            with pd.ExcelWriter(file_path, engine='openpyxl', mode='a') as writer:
                # Get the workbook and the sheet
                workbook = writer.book
                sheet = writer.sheets['sheet_1']
                
                # Set the column widths (you can adjust the width values as needed)
                column_widths = {
                    'A': 10,
                    'B': 15,
                    'C': 10,
                    'D': 15,
                    'E': 15,
                    'F': 25,
                    'G': 10,
                    'H': 10,
                    'I': 15
                }

                for col, width in column_widths.items():
                    sheet.column_dimensions[col].width = width

                # Hide the first two columns (A and B)
                sheet.column_dimensions['A'].hidden = True
                sheet.column_dimensions['B'].hidden = True

        else:
            logging.error(f"No enrollments found for subject option: {options_name} in subject type: {subject_type}")
            
    total_students = sum(summary_data.values())
    total_options = len(summary_data)
    
    summary_logging.info(f"Subject Type: {subject_type}")
    for option_name, student_count in summary_data.items():
        summary_logging.info(f"Option Name: {option_name}, Students: {student_count}")

    summary_logging.info(f"Total Option Types: {total_options}, Total Students: {total_students}\n")
    
    summary_handler.flush()

print("Processing complete.")
